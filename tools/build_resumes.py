from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import os

OUT = '.'
BLUE = RGBColor(48, 65, 171)
LINK_BLUE = RGBColor(0, 88, 174)
TEXT = RGBColor(27, 27, 30)
SECOND = RGBColor(61, 67, 75)
FNT = 'Liberation Sans'

def hyperlink(paragraph, text, url, size=9, color=LINK_BLUE, bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    font = OxmlElement('w:rFonts')
    font.set(qn('w:ascii'), FNT)
    font.set(qn('w:hAnsi'), FNT)
    pr.append(font)
    if bold:
        pr.append(OxmlElement('w:b'))
    col = OxmlElement('w:color')
    col.set(qn('w:val'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pr.append(col)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(size*2)))
    pr.append(sz)
    r.append(pr)
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)

def add_runs(p, parts, font_size=8.55, color=TEXT):
    for txt, bold in parts:
        r = p.add_run(txt)
        r.bold = bold
        r.font.name = FNT
        r.font.size = Pt(font_size)
        r.font.color.rgb = color
    return p

def border_bottom(p, color='3041AB'):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), '7')
    b.set(qn('w:space'), '3')
    b.set(qn('w:color'), color)
    borders.append(b)
    pPr.append(borders)

def add_heading(doc, txt, before=8, after=11):
    p = doc.add_paragraph(style='SectionHead')
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt)
    r.bold = True
    r.font.name = FNT
    r.font.size = Pt(10.25)
    r.font.color.rgb = BLUE
    border_bottom(p)
    return p

def body(doc, parts, after=1.6, before=0, size=8.65, leading=11.35):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = Pt(leading)
    return add_runs(p, parts, size)

def bullet(doc, parts, after=5.6):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.left_indent = Pt(10)
    f.first_line_indent = Pt(-10)
    f.space_before = Pt(0)
    f.space_after = Pt(after)
    f.line_spacing = Pt(12.55)
    add_runs(p, [('• ', False)] + parts, 8.53)

def compact_title(doc, parts, after=1, before=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(11)
    add_runs(p, parts, 9)

def split_bold(text, fragments):
    parts = []
    remaining = text
    for target in fragments:
        ix = remaining.find(target)
        if ix < 0:
            raise ValueError(f'Missing bold snippet: {target}')
        if ix:
            parts.append((remaining[:ix], False))
        parts.append((target, True))
        remaining = remaining[ix+len(target):]
    if remaining:
        parts.append((remaining, False))
    return parts

rows_common = [
('HPL (Haldia Petrochemicals): Imported 5,000+ instrument tags for JB wiring; configured device panels and cables, updated instrument index data and checked wiring connections. Mapped and imported specification, process and line data; maintained datasheet and JB revisions.', ['HPL (Haldia Petrochemicals):', '5,000+ instrument tags']),
('MPL (Mundra Petrochemicals): Updated 200+ instrument tags, created loops and corrected reference panel/cable assignments. Converted AutoCAD hook-ups to .sym, assigned tags and updated drawings and BOMs.', ['MPL (Mundra Petrochemicals):', '200+ instrument tags', '.sym']),
('IOCL PBR: Imported 600 JBs for DCS, ESD and F&G systems; configured panels, terminals and cables. Prepared hook-up items and supported MCC and marshalling rack wiring.', ['IOCL PBR:', '600 JBs', 'MCC and marshalling rack wiring']),
('RIL VCM-021 & DMD-121: Maintained Spec Data Dictionary settings and datasheet alignment; executed wiring for 700+ tags. Imported 430 .ISF files for the VCM refrigeration package in batches.', ['RIL VCM-021 & DMD-121:', 'Spec Data Dictionary', '700+ tags', '430 .ISF files']),
('RIL EDC: Configured JB, MCC and MR wiring for 1,000+ tags; checked terminal connections, cable assignments and signal levels in SPI.', ['RIL EDC:', '1,000+ tags']),
('HZL & ACME: Configured branch/main cables to approved project requirements (FR/FRLS, IS/NIS); generated cable drum schedules and JB wiring reports.', ['HZL & ACME:', 'FR/FRLS, IS/NIS']),
]
summary = 'Instrumentation Engineer with 2 years 4 months of SmartPlant Instrumentation (SPI / INtools) experience supporting oil & gas and petrochemical projects, plus a one-year graduate apprenticeship at BPCL Mumbai Refinery. Skilled in instrument index management, datasheets, bulk imports, wiring configurations and hook-up deliverables, with hands-on exposure to instrument maintenance and calibration.'

def generate(version):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.2677)
    sec.page_height = Inches(11.6929)
    sec.top_margin = Inches(.56)
    sec.bottom_margin = Inches(.53)
    sec.left_margin = Inches(.55)
    sec.right_margin = Inches(.55)
    sec.header_distance = Inches(.18)
    sec.footer_distance = Inches(.18)

    normal = doc.styles['Normal']
    normal.font.name = FNT
    normal.font.size = Pt(8.65)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(11.35)
    hstyle = doc.styles.add_style('SectionHead', WD_STYLE_TYPE.PARAGRAPH)
    hstyle.base_style = normal

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(22.4)
    r = p.add_run('SALMAN PAWASKAR')
    r.bold = True
    r.font.name = FNT
    r.font.size = Pt(20.9)
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(11.3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.15), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('Instrumentation Engineer | SPI / INtools')
    r.font.name = FNT
    r.font.size = Pt(9.5)
    r.font.color.rgb = BLUE
    p.add_run('\t')
    hyperlink(p, 'LinkedIn profile', 'https://www.linkedin.com/in/salman-pawaskar-6039801a1', 8.5, SECOND, False)

    contact = ('Mumbai, India  |  +91 8356829155  |  pawaskarsalman47@gmail.com' if version == 'India'
               else 'Mumbai, India (Open to GCC relocation)  |  +91 8356829155  |  pawaskarsalman47@gmail.com')
    body(doc, [(contact, False)], after=7.5, size=8.8, leading=11.2)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12.5)
    p.paragraph_format.line_spacing = Pt(16)
    r = p.add_run('PORTFOLIO')
    r.bold = True
    r.font.name = FNT
    r.font.size = Pt(9.8)
    r.font.color.rgb = RGBColor(0, 0, 0)
    r = p.add_run('     ')
    r.font.size = Pt(10)
    hyperlink(p, 'www.salmanpawaskar.com', 'https://salmanpawaskar.com', 11.5, LINK_BLUE, True)

    add_heading(doc, 'PROFESSIONAL SUMMARY', before=0, after=11)
    body(doc, split_bold(summary, ['2 years 4 months of SmartPlant Instrumentation (SPI / INtools) experience']), after=7.0, size=8.8, leading=12.8)

    add_heading(doc, 'TECHNICAL SKILLS', before=4, after=11)
    body(doc, [('SPI: ', True), ('Instrument Index, Specifications, Import Utility, Browser, Engineering Data Editor (EDE)', False)], after=2, leading=12, size=8.65)
    body(doc, [('Deliverables: ', True), ('Instrument Datasheets, Process & Line Data, Hook-Up Drawings, Bills of Materials, Cable Schedules, Loop Drawing Verification', False)], after=2, leading=12, size=8.65)
    body(doc, [('Wiring: ', True), ('Junction Box (JB), Motor Control Centre (MCC), Marshalling Rack (MR)', False)], after=2, leading=12, size=8.65)
    if version == 'India':
        body(doc, [('Data & Reporting: ', True), ('Microsoft Excel, SPI Query Builder, PSR Reports', False)], after=2, leading=12, size=8.65)

    add_heading(doc, 'PROFESSIONAL EXPERIENCE', before=8, after=11)
    compact_title(doc, [('Krishna Global Services Pvt. Ltd. | Instrumentation Engineer', True)], after=1.7, before=0)
    body(doc, [('Deployed at ', False), ('Thyssenkrupp Uhde India Pvt. Ltd.', True), (', Mumbai', False)], after=0.3, size=8.55, leading=10.9)
    body(doc, [('May 2024 - September 2026', False)], after=4, size=8.45, leading=10.8)

    for text, spans in rows_common:
        bullet(doc, split_bold(text, spans), after=5.6)
    if version == 'India':
        last = 'Across projects: Used Browser, EDE and Query Builder for bulk updates; cross-checked SPI extracts in Excel for missing connections and terminal mismatches. Created hook-up variants for 100+ tags to meet the 20-tags-per-page requirement.'
    else:
        last = 'Across projects: Used Browser, EDE and Query Builder for bulk updates and checks. Created hook-up variants for 100+ tags to meet the 20-tags-per-page requirement.'
    bullet(doc, split_bold(last, ['Across projects:', '100+ tags', '20-tags-per-page']), after=6.5)

    compact_title(doc, [('Bharat Petroleum Corporation Limited | Graduate Apprentice Trainee', True)], after=1.7, before=0.5)
    body(doc, [('Mumbai Refinery | March 2023 - March 2024', False)], after=3.8, size=8.45, leading=10.8)
    bpcl1 = 'Performed control valve maintenance, stroke checks, commissioning and calibration at CDU-3; calibrated flow transmitters and gained exposure to nucleonic gauge calibration at CCR.'
    bpcl2 = 'Participated in CCU, FCCU, GTU and CCR shutdown activities and refinery Fire & Safety Training. Checked MOC task checklists, LIMS reports and COMOS workflows.'
    bullet(doc, split_bold(bpcl1, ['control valve maintenance, stroke checks, commissioning and calibration']), after=5.6)
    bullet(doc, split_bold(bpcl2, ['CCU, FCCU, GTU and CCR shutdown activities']), after=5.6)

    add_heading(doc, 'EDUCATION', before=7, after=11)
    body(doc, [('B.E. - Instrumentation Engineering | 2022', True)], after=0, size=8.8, leading=10.9)
    body(doc, [('MCT’s Rajiv Gandhi Institute of Technology, University of Mumbai', False)], after=5.4, size=8.5, leading=10.9)
    body(doc, [('Languages: ', True), ('English, Hindi, Marathi', False)], after=0, size=8.5, leading=10.9)

    core = doc.core_properties
    core.title = f'Salman Pawaskar | {version} Resume'
    core.subject = 'Instrumentation Engineering | SPI / INtools'
    core.author = 'Salman Pawaskar'
    core.keywords = 'Instrumentation Engineer, SPI, INtools, SmartPlant, Instrumentation'
    docx = f'{OUT}/Salman_Pawaskar_{version}_Resume_Updated.docx'
    doc.save(docx)
    print(docx)

for v in ('India', 'GCC'):
    generate(v)
